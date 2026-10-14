"""Parcel and valuation data from the Johnson County AIMS parcel export.

The county export has one row per map polygon, not per tax parcel: a condo or
quadraplex unit appears once for its own unit polygon and again for the
shared common-area polygon, with the same TaxPropID and values on both. Import
collapses the rows to one record per TaxPropID, keeping the unit's own polygon
and recording the common-area acreage separately, so totals count each
parcel once.

Each year is stored as parcels/<year>.csv, checked in so year-over-year
comparisons can diff them.
"""

import csv
import re
import statistics
from collections import Counter, defaultdict
from pathlib import Path

from pvcouncil.config import PROJECT_ROOT

PARCELS_DIR = PROJECT_ROOT / "parcels"

# Output column -> source column in the AIMS export
SOURCE_COLUMNS = {
    "parcel_id": "TaxPropID",
    "kupn": "TaxKUPN",
    "address": "SALine1",
    "owner1": "Own1Name",
    "owner2": "Own2Name",
    "mail_name": "MailName",
    "mail_line1": "OALine1",
    "mail_line2": "OALine2",
    "class": "GenClass",
    "land_use": "GenLandUs1",
    "function": "LBCSFunDsc",
    "structure": "LBCSStrDsc",
    "zoning": "Zoning",
    "tax_unit": "TaxUnit",
    "dwelling_units": "DwellUnits",
    "year_built": "YearBuilt",
    "style": "ResStylDsc",
    "grade": "GradeDesc",
    "cdu": "CDUDesc",
    "condition": "CondDesc",
    "acres": "AreaCAcre",
    "lot_sqft": "AreaCSqFt",
    "land_value": "LandValue",
    "improvement_value": "ImprValue",
    "total_value": "TotalValue",
    "assessed_land": "AssLandVal",
    "assessed_improvement": "AssImprVal",
    "assessed_total": "AssTotVal",
    "assessment_pct_land": "AssLandPct",
    "assessment_pct_improvement": "AssImprPct",
    "polygon_type": "PrPlTypDsc",
    "legal": "LegalDesc",
    "url": "PropIDURL",
}
PLAT_COLUMNS = ["subdivision_code", "subdivision", "plat", "block", "lot"]
# Owner names and mailing addresses go to a separate, gitignored
# parcels/<year>-owners.csv so they stay local and out of the repo.
OWNER_COLUMNS = ["owner1", "owner2", "mail_name", "mail_line1", "mail_line2"]
# Flags derived from the owner data that say nothing about who the owner is
OWNER_FLAG_COLUMNS = ["corp_or_llc_owned", "owner_out_of_state", "owner_mail_state"]
COLUMNS = ([c for c in SOURCE_COLUMNS if c not in OWNER_COLUMNS] + PLAT_COLUMNS
           + ["common_area_acres", "polygons"] + OWNER_FLAG_COLUMNS)

INT_COLUMNS = {"dwelling_units", "year_built", "lot_sqft", "land_value", "improvement_value",
               "total_value", "assessed_land", "assessed_improvement", "assessed_total", "polygons",
               "corp_or_llc_owned", "owner_out_of_state"}
FLOAT_COLUMNS = {"acres", "common_area_acres", "assessment_pct_land", "assessment_pct_improvement"}

# A single-family house, as opposed to a condo, townhouse, or duplex unit
SFR = "Single family residence"


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------

def _pick_sheet(wb, sheet: str | None):
    """The named sheet, else the sheet with the most rows (the unfiltered export)."""
    if sheet:
        return wb[sheet]
    return max(wb.worksheets, key=lambda ws: ws.max_row)


def import_xlsx(xlsx_path: Path, year: int, sheet: str | None = None) -> dict:
    """Read an AIMS parcel export and write parcels/<year>.csv. Returns stats."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = _pick_sheet(wb, sheet)
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    missing = [c for c in [*SOURCE_COLUMNS.values(), "TaxBlock", "TaxLot"] if c not in header]
    if missing:
        raise ValueError(f"sheet {ws.title!r} is missing columns: {', '.join(missing)}")

    groups: dict[str, list[dict]] = defaultdict(list)
    skipped = 0
    for raw in rows:
        r = dict(zip(header, raw))
        # Hand-added summary rows (averages, medians) have no parcel ID
        if not r.get("TaxPropID"):
            skipped += int(any(v is not None for v in raw))
            continue
        groups[r["TaxPropID"]].append(r)

    records, conflicts = [], []
    for pid, polys in groups.items():
        own = [p for p in polys if p["GeoPropID"] == pid]
        main = own[0] if own else polys[0]
        if len({p["TotalValue"] for p in polys}) > 1:
            conflicts.append(pid)
        rec = {k: main[src] for k, src in SOURCE_COLUMNS.items()}
        common = [p["AreaCAcre"] for p in polys if p is not main and p["PrPlTypDsc"] == "Common Interest Property Polygon"]
        rec["common_area_acres"] = common[0] if common else None
        rec["polygons"] = len(polys)
        rec.update(_plat_fields(pid, main))
        rec = _clean(rec)
        rec.update(_owner_flags(rec))
        records.append(rec)
    _name_subdivisions(records)

    records.sort(key=lambda r: r["parcel_id"])
    PARCELS_DIR.mkdir(exist_ok=True)
    out = PARCELS_DIR / f"{year}.csv"
    for path, fields in ((out, COLUMNS), (owners_path(year), ["parcel_id", *OWNER_COLUMNS])):
        with path.open("w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields, lineterminator="\n", extrasaction="ignore")
            w.writeheader()
            w.writerows(records)

    return {"sheet": ws.title, "rows": sum(len(p) for p in groups.values()),
            "parcels": len(records), "skipped": skipped, "conflicts": conflicts, "path": out}


# Where the plat name ends in a legal description: the lot/block/unit, a
# bearing or partial-lot clause ("NWLY 8 FT", "W1/2", "BG SW CR"), or a comma
# or parenthetical.
_PLAT_END = re.compile(
    r"\s*(?:[,(]"
    r"|\b(?:LOTS?|LTS?|BLKS?|BLOCKS?|BLDGS?|BUILDING|UNITS?|TRACTS?|TR|PTS?|PTLTS?|ALL|EX|TRI|BG|BEG|CERTIFICATE|[NSEW]{1,2}LY)\b"
    r"|\b(?:PT)?LTS?\d"
    r"|\b(?:[NSEW]{1,2}|NORTH|SOUTH|EAST|WEST)\s*[\d.]+"
    r"|\b[NSEW]{1,2}\s*\d/\d).*$")


def _plat_fields(pid: str, row: dict) -> dict:
    """Subdivision code, plat name, block, and lot for a platted parcel.

    Platted parcel IDs read OP<subdivision><block> <lot>, e.g. OP55000002 0030
    is subdivision 5500 (Prairie Village), block 2, lot 30. One subdivision
    code can span several plats (West Riding and West Riding 2nd Plat), so
    the plat name parsed from the legal description is kept as well.
    Unplatted acreage (OF...) has none of these.
    """
    if not pid.startswith("OP"):
        return dict.fromkeys(PLAT_COLUMNS)
    # A few descriptions lead with the block ("BLOCK 1 MISSION VILLAGE LT 91 ..."),
    # and replats name the lots they redrew ("RESURVEY OF LOTS 15 AND 16 RIDGE VIEW LT 5 ...").
    legal = re.sub(r"^BLOCK \w+ ", "", row["LegalDesc"] or "")
    m = re.match(r"(?:REPLAT|RESURVEY) OF (?:PART OF )?LOTS? \d+(?:\s*(?:AND|&)\s*\d+)*\s+", legal)
    prefix, legal = (m.group(0), legal[m.end():]) if m else ("", legal)
    plat = (prefix + _PLAT_END.sub("", legal)).strip(" ,.-")
    plat = re.sub(r"\s+OF$", "", plat) or None
    return {"subdivision_code": pid[2:6], "subdivision": None, "plat": plat,
            "block": (row["TaxBlock"] or "").lstrip("0") or None,
            "lot": (row["TaxLot"] or "").lstrip("0") or None}


def _name_subdivisions(records: list[dict]) -> None:
    """Name each subdivision code after its most common plat name."""
    names = defaultdict(Counter)
    for r in records:
        if r["subdivision_code"] and r["plat"]:
            names[r["subdivision_code"]][r["plat"]] += 1
    for r in records:
        if r["subdivision_code"] in names:
            r["subdivision"] = names[r["subdivision_code"]].most_common(1)[0][0]


_LLC = re.compile(r"\bL\.?\s?L\.?\s?C\b|\bL\.?C\b")
# A bare "CO" is left out because in these names it is nearly always
# "CO-TRUSTEE", not "& Co."
_CORP = re.compile(r"\b(?:INC|INCORPORATED|CORP|CORPORATION|COMPANY|LTD)\b")
# Partnerships: LP, L.P., LLP, LLLP, FLP (family limited partnership),
# "... PARTNERS", "... LIMITED PARTNERSHIP". There is no space inside the
# initials, so a middle-initial pair like "JOHN L P SMITH" doesn't match.
_PARTNERSHIP = re.compile(r"\b(?:L\.?P|L\.?L\.?P|L\.?L\.?L\.?P|F\.?L\.?P)\b|\bPARTNERS(?:HIP)?\b")


def _is_corp_or_llc(name: str | None) -> bool:
    if not name:
        return False
    if _LLC.search(name):
        return True
    # "CORP, LILLIAN I" is a surname. A trust company acting as trustee
    # holds the property for a person's trust.
    if re.match(r"[A-Z'-]+, ", name) or re.search(r"\b(?:TRUSTEE|TTEE)\b", name):
        return False
    return bool(_CORP.search(name) or _PARTNERSHIP.search(name))


def _owner_flags(rec: dict) -> dict:
    """Business-entity ownership and the owner's mailing state, from the owner fields.

    corp_or_llc_owned covers corporations, LLCs, and partnerships. Only the
    owner names count toward it. The mailing
    name is often a bank or tax service paying the bill for the owner. The
    state is the two-letter code before the ZIP on the last mailing line
    ("KANSAS CITY, MO 64114"). A foreign address has no state code and is
    recorded as out of state with no state. A missing address is unknown,
    and both flags are left blank.
    """
    corp = int(_is_corp_or_llc(rec["owner1"]) or _is_corp_or_llc(rec["owner2"]))
    line = rec["mail_line2"]
    if not line:
        return {"corp_or_llc_owned": corp, "owner_out_of_state": None, "owner_mail_state": None}
    m = re.search(r"\b([A-Z]{2})\s+\d{5}(?:-\d{4})?$", line)
    state = m.group(1) if m else None
    return {"corp_or_llc_owned": corp, "owner_out_of_state": int(state != "KS"), "owner_mail_state": state}


def _clean(rec: dict) -> dict:
    for k, v in rec.items():
        if isinstance(v, str):
            v = re.sub(r"\s+", " ", v).strip() or None
        elif k in INT_COLUMNS and v is not None:
            v = int(round(v))
        elif k in FLOAT_COLUMNS and v is not None:
            v = round(float(v), 4)
        rec[k] = v
    return rec


# ---------------------------------------------------------------------------
# Loading and querying
# ---------------------------------------------------------------------------

def available_years() -> list[int]:
    return sorted(int(p.stem) for p in PARCELS_DIR.glob("*.csv") if p.stem.isdigit())


def load(year: int | None = None) -> list[dict]:
    """Parcel records for a year (default: the latest imported)."""
    years = available_years()
    if not years:
        raise FileNotFoundError("No parcel data yet — run `pvc import-parcels` first.")
    year = year or years[-1]
    path = PARCELS_DIR / f"{year}.csv"
    if not path.exists():
        raise FileNotFoundError(f"No parcel data for {year} (have: {', '.join(map(str, years))}).")
    with path.open(newline="") as f:
        records = [_typed(r) for r in csv.DictReader(f)]
    owner_info = {}
    if owners_path(year).exists():
        with owners_path(year).open(newline="") as f:
            owner_info = {r["parcel_id"]: _typed(r) for r in csv.DictReader(f)}
    blank = dict.fromkeys(OWNER_COLUMNS)
    for r in records:
        r.update({k: owner_info.get(r["parcel_id"], blank)[k] for k in OWNER_COLUMNS})
    return records


def owners_path(year: int) -> Path:
    return PARCELS_DIR / f"{year}-owners.csv"


def has_owners(parcels: list[dict]) -> bool:
    return any(p["owner1"] for p in parcels)


def _typed(r: dict) -> dict:
    out = {}
    for k, v in r.items():
        if v == "":
            v = None
        elif k in INT_COLUMNS:
            v = int(v)
        elif k in FLOAT_COLUMNS:
            v = float(v)
        out[k] = v
    return out


def _addr_key(s: str | None) -> str:
    return re.sub(r"[^A-Z0-9]", "", (s or "").upper())


def mails_elsewhere(p: dict) -> bool:
    """Tax bill goes somewhere other than the property itself (a rough non-owner-occupied signal)."""
    return bool(p["mail_line1"]) and not _addr_key(p["mail_line1"]).startswith(_addr_key(p["address"]).split("#")[0])


def owners(p: dict) -> str:
    return " & ".join(o for o in (p["owner1"], p["owner2"]) if o)


def query(parcels: list[dict], *, address=None, owner=None, zoning=None, cls=None, use=None,
          tax_unit=None, subdivision=None, block=None, built_after=None, built_before=None,
          min_value=None, max_value=None, mails_elsewhere_only=False, corp_or_llc=False,
          out_of_state=False, mail_state=None) -> list[dict]:
    def match(p):
        if corp_or_llc and not p["corp_or_llc_owned"]:
            return False
        if out_of_state and not p["owner_out_of_state"]:
            return False
        if mail_state and (p["owner_mail_state"] or "") != mail_state.upper():
            return False
        if address and address.upper() not in (p["address"] or ""):
            return False
        if subdivision and subdivision.upper() not in f"{p['subdivision'] or ''}|{p['plat'] or ''}" \
                and subdivision != p["subdivision_code"]:
            return False
        if block and (p["block"] or "").upper() != block.upper().lstrip("0"):
            return False
        if owner and owner.upper() not in " ".join(filter(None, (p["owner1"], p["owner2"], p["mail_name"]))).upper():
            return False
        if zoning and (p["zoning"] or "").upper() != zoning.upper():
            return False
        if cls and (p["class"] or "").lower() != cls.lower():
            return False
        if use and use.lower() not in (p["function"] or "").lower():
            return False
        if tax_unit and p["tax_unit"] != tax_unit:
            return False
        if built_after and (p["year_built"] or 0) < built_after:
            return False
        if built_before and (not p["year_built"] or p["year_built"] > built_before):
            return False
        if min_value is not None and (p["total_value"] or 0) < min_value:
            return False
        if max_value is not None and (p["total_value"] or 0) > max_value:
            return False
        if mails_elsewhere_only and not mails_elsewhere(p):
            return False
        return True
    return [p for p in parcels if match(p)]


# ---------------------------------------------------------------------------
# Summaries
# ---------------------------------------------------------------------------

def _pct(values: list[int], q: float) -> int:
    return values[min(len(values) - 1, int(q * len(values)))]


def value_stats(parcels: list[dict]) -> dict:
    vals = sorted(p["total_value"] for p in parcels if p["total_value"])
    if not vals:
        return {"n": 0}
    return {"n": len(vals), "median": statistics.median(vals), "p10": _pct(vals, 0.1),
            "p90": _pct(vals, 0.9), "min": vals[0], "max": vals[-1]}


def totals(parcels: list[dict]) -> dict:
    return {"parcels": len(parcels),
            "appraised": sum(p["total_value"] or 0 for p in parcels),
            "assessed": sum(p["assessed_total"] or 0 for p in parcels),
            "dwelling_units": sum(p["dwelling_units"] or 0 for p in parcels)}


def group_totals(parcels: list[dict], key) -> list[tuple[str, dict]]:
    groups = defaultdict(list)
    for p in parcels:
        groups[key(p)].append(p)
    return sorted(((k, totals(v)) for k, v in groups.items()), key=lambda kv: -kv[1]["assessed"])


def budget_assessed_valuation(year: int) -> tuple[int, float] | None:
    """Assessed valuation the budget used for this valuation year, from budgets.json if present.

    A January 1 valuation sets the levy for the following budget year.
    """
    try:
        from pvcouncil.budget import BUDGETS_PATH, budget_years
        if not BUDGETS_PATH.exists():
            return None
        rec = budget_years().get(year + 1)
    except Exception:
        return None
    v = rec and rec["headline"].get("assessed_valuation", {}).get("value")
    return (year + 1, v) if v else None


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------

def money(v: float) -> str:
    if abs(v) >= 1e9:
        return f"${v / 1e9:,.2f}B"
    if abs(v) >= 1e6:
        return f"${v / 1e6:,.1f}M"
    return f"${v:,.0f}"


def render_list(parcels: list[dict], n: int = 50) -> str:
    out = []
    for p in parcels[:n] if n else parcels:
        built = p["year_built"] or "—"
        out.append(f"{p['address'] or '(no address)':<28} {money(p['total_value'] or 0):>10}  "
                   f"assessed {money(p['assessed_total'] or 0):>9}  {p['zoning'] or '—':<6} {built!s:<5} "
                   f"{(p['function'] or '')[:28]:<28} {owners(p)}")
    if n and len(parcels) > n:
        out.append(f"... {len(parcels) - n} more (use -n 0 for all)")
    return "\n".join(out)


def _plat_line(p: dict) -> str:
    if not p["subdivision_code"]:
        return "unplatted acreage"
    name = p["plat"] or p["subdivision"]
    if p["subdivision"] and name != p["subdivision"]:
        name += f" (part of {p['subdivision']})"
    return f"{name}, block {p['block'] or '—'}, lot {p['lot'] or '—'}, subdivision code {p['subdivision_code']}"


def subdivision_rows(parcels: list[dict]) -> list[dict]:
    """Per-subdivision totals and medians, largest first; unplatted acreage last."""
    groups = defaultdict(list)
    for p in parcels:
        groups[(p["subdivision_code"], p["subdivision"])].append(p)
    rows = []
    for (code, name), ps in groups.items():
        homes = [p for p in ps if p["function"] == SFR and (p["total_value"] or 0) > 0]
        years = [p["year_built"] for p in ps if p["year_built"]]
        lots = [p["lot_sqft"] for p in homes if p["lot_sqft"]]
        rows.append({"code": code, "name": name or "(unplatted acreage)", **totals(ps),
                     "homes": len(homes),
                     "median_home": statistics.median(p["total_value"] for p in homes) if homes else None,
                     "median_year": int(statistics.median(years)) if years else None,
                     "median_lot": statistics.median(lots) if lots else None})
    return sorted(rows, key=lambda r: (r["code"] is None, -r["parcels"]))


def render_detail(p: dict) -> str:
    lines = [f"{p['address']}  ({p['parcel_id']})"]
    if p["owner1"]:
        lines += [f"  Owner:       {owners(p)}",
                  f"  Mailing:     {', '.join(filter(None, (p['mail_name'], p['mail_line1'], p['mail_line2'])))}"]
    if p["owner_out_of_state"] is not None:
        state = p["owner_mail_state"] or "outside the US"
        lines.append(f"  Ownership:   {'corporation/LLC/partnership-owned' if p['corp_or_llc_owned'] else 'not owned by a corporation, LLC, or partnership'}, "
                     f"owner mails from {state}{' (out of state)' if p['owner_out_of_state'] else ''}")
    lines += [f"  Use:         {p['function']} ({p['class']}), zoned {p['zoning']}, tax unit {p['tax_unit']}",
              f"  Built:       {p['year_built'] or '—'}  {p['style'] or ''}  grade {p['grade'] or '—'}, "
              f"CDU {p['cdu'] or '—'}, condition {p['condition'] or '—'}",
              f"  Lot:         {p['acres']} acres ({p['lot_sqft']:,} sq ft)" if p["lot_sqft"] else "  Lot:         —",
              f"  Appraised:   {money(p['total_value'] or 0)} (land {money(p['land_value'] or 0)}, "
              f"improvements {money(p['improvement_value'] or 0)})",
              f"  Assessed:    {money(p['assessed_total'] or 0)} "
              f"({p['assessment_pct_land']}% land / {p['assessment_pct_improvement']}% improvements)",
              f"  Plat:        {_plat_line(p)}",
              f"  Legal:       {p['legal']}",
              f"  County:      {p['url']}"]
    return "\n".join(lines)


def render_summary(parcels: list[dict]) -> str:
    t = totals(parcels)
    s = value_stats(parcels)
    out = [f"{t['parcels']:,} parcels, {t['dwelling_units']:,} dwelling units",
           f"Appraised {money(t['appraised'])}, assessed {money(t['assessed'])}"]
    if s["n"]:
        out.append(f"Value: median {money(s['median'])}, 10th pct {money(s['p10'])}, "
                   f"90th pct {money(s['p90'])}, max {money(s['max'])}")
    out += ["", "By class:"]
    for k, g in group_totals(parcels, lambda p: p["class"] or "—"):
        out.append(f"  {k:<16} {g['parcels']:>6,}  appraised {money(g['appraised']):>9}  assessed {money(g['assessed']):>9}")
    out += ["", "By zoning:"]
    for k, g in group_totals(parcels, lambda p: p["zoning"] or "—")[:12]:
        out.append(f"  {k:<16} {g['parcels']:>6,}  appraised {money(g['appraised']):>9}  assessed {money(g['assessed']):>9}")
    subs = subdivision_rows(parcels)
    if len(subs) > 1:
        out += ["", "By subdivision:"]
        for r in subs[:15]:
            home = f"median home {money(r['median_home'])}" if r["median_home"] else ""
            out.append(f"  {r['name'][:30]:<30} {r['parcels']:>6,}  assessed {money(r['assessed']):>9}  {home}")
        if len(subs) > 15:
            out.append(f"  ... {len(subs) - 15} more")
    return "\n".join(out)


def _table(header: list[str], rows: list[list], right: set[int] | None = None) -> list[str]:
    right = right if right is not None else set(range(1, len(header)))
    lines = ["| " + " | ".join(header) + " |",
             "|" + "|".join(" ---: " if i in right else " --- " for i in range(len(header))) + "|"]
    lines += ["| " + " | ".join(str(c) for c in r) + " |" for r in rows]
    return lines


def render_markdown(year: int | None = None) -> str:
    """Property and valuation report for one year."""
    year = year or available_years()[-1]
    ps = load(year)
    t = totals(ps)
    taxable = [p for p in ps if (p["assessed_total"] or 0) > 0]
    sfr = [p for p in ps if p["function"] == SFR and p["class"] == "Residential" and (p["total_value"] or 0) > 0]
    ss = value_stats(sfr)

    lines = [f"# Prairie Village Property and Valuation, {year}", "",
             f"Built from the Johnson County AIMS parcel export for {year} "
             f"(`parcels/{year}.csv`, regenerate with `pvc parcels --markdown`). Values are the "
             "county's January 1 appraisals. Assessed value is the share of appraised value "
             "that is taxed: 11.5% for residential and 25% for commercial property. Condo and "
             "townhouse units show up twice in the county export, once for the unit and once "
             "for the shared common area. Each is counted once here.", "",
             "## Overview", ""]
    lines += _table(["", ""], [
        ["Tax parcels", f"{t['parcels']:,}"],
        ["Dwelling units", f"{t['dwelling_units']:,}"],
        ["Appraised value", money(t["appraised"])],
        ["Assessed value", money(t["assessed"])],
        ["Taxable parcels", f"{len(taxable):,}"],
        ["Single-family homes", f"{ss['n']:,}"],
        ["Median single-family appraisal", money(ss["median"])],
    ])

    bav = budget_assessed_valuation(year)
    if bav:
        by, v = bav
        diff = (t["assessed"] - v) / v * 100
        lines += ["", f"The {by} budget uses an assessed valuation of {money(v)}. "
                  f"The parcels here add up to {money(t['assessed'])} ({diff:+.1f}%). "
                  "The budget figure also covers personal and state-assessed utility property, "
                  "which have no parcels."]

    lines += ["", "## By property class", ""]
    lines += _table(["Class", "Parcels", "Appraised", "Assessed", "Share of assessed"],
                    [[k, f"{g['parcels']:,}", money(g["appraised"]), money(g["assessed"]),
                      f"{g['assessed'] / t['assessed'] * 100:.1f}%"]
                     for k, g in group_totals(ps, lambda p: p["class"] or "—")])

    lines += ["", "## By zoning district", ""]
    zrows = []
    for k, g in group_totals(ps, lambda p: p["zoning"] or "(none)"):
        zs = value_stats([p for p in sfr if (p["zoning"] or "(none)") == k])
        zrows.append([k, f"{g['parcels']:,}", f"{g['dwelling_units']:,}", money(g["appraised"]),
                      money(g["assessed"]), money(zs["median"]) if zs["n"] >= 5 else "—"])
    lines += _table(["Zoning", "Parcels", "Dwelling units", "Appraised", "Assessed", "Median single-family"], zrows)

    subs = subdivision_rows(ps)
    lines += ["", "## By subdivision", "",
              f"{sum(1 for r in subs if r['code'])} subdivisions, grouped by the subdivision code "
              "in the county parcel ID. A code can span more than one plat (for example, "
              "West Riding and West Riding 2nd Plat). Each subdivision is named after its "
              "most common plat. Unplatted acreage is described by survey instead of a lot "
              "and block, and is listed last. Medians cover single-family homes only.", ""]
    lines += _table(["Subdivision", "Code", "Parcels", "Homes", "Median home", "Median year built",
                     "Median lot (sq ft)", "Appraised"],
                    [[r["name"], r["code"] or "—", f"{r['parcels']:,}", f"{r['homes']:,}",
                      money(r["median_home"]) if r["median_home"] else "—", r["median_year"] or "—",
                      f"{r['median_lot']:,.0f}" if r["median_lot"] else "—", money(r["appraised"])]
                     for r in subs], right={2, 3, 4, 5, 6, 7})

    lines += ["", "## Housing types", ""]
    res = [p for p in ps if p["class"] == "Residential"]
    hrows = []
    for k, n in Counter(p["function"] for p in res).most_common():
        grp = [p for p in res if p["function"] == k]
        vs = value_stats(grp)
        hrows.append([k or "—", f"{n:,}", f"{sum(p['dwelling_units'] or 0 for p in grp):,}",
                      money(vs["median"]) if vs["n"] else "—"])
    lines += _table(["Use", "Parcels", "Dwelling units", "Median appraisal"], hrows)

    lines += ["", "## Single-family homes", "",
              f"{ss['n']:,} homes. Median appraisal {money(ss['median'])}. "
              f"10% are under {money(ss['p10'])} and 10% are over {money(ss['p90'])}. "
              f"The highest is {money(ss['max'])}. "
              f"Land is a median {statistics.median(p['land_value'] / p['total_value'] for p in sfr) * 100:.0f}% "
              "of a home's appraised value.", "", "**By decade built:**", ""]
    decades = defaultdict(list)
    for p in sfr:
        if p["year_built"]:
            decades[p["year_built"] // 10 * 10].append(p)
    drows = []
    for d in sorted(decades):
        if len(decades[d]) < 5:
            continue
        vs = value_stats(decades[d])
        drows.append([f"{d}s", f"{vs['n']:,}", money(vs["median"]),
                      f"{statistics.median(p['lot_sqft'] for p in decades[d] if p['lot_sqft']):,.0f}"])
    lines += _table(["Built", "Homes", "Median appraisal", "Median lot (sq ft)"], drows)
    lines += ["", "Decades with fewer than five homes are left out.", ""]

    recent = [p for p in sfr if (p["year_built"] or 0) >= year - 10]
    if recent:
        lines += [f"**Built {year - 10}–{year - 1}:**", ""]
        by_year = defaultdict(list)
        for p in recent:
            by_year[p["year_built"]].append(p)
        lines += _table(["Year", "Homes", "Median appraisal"],
                        [[y, len(by_year[y]), money(value_stats(by_year[y])["median"])] for y in sorted(by_year)])
        lines += ["", "Almost every lot in Prairie Village was built on by the 1970s. A recent "
                  "year built usually means an older house was torn down and replaced."]

    # Owner names stay out of the report: only citywide shares, from the
    # corp_or_llc_owned / owner_out_of_state flags, plus the mailed-elsewhere
    # share when the local owners file is present.
    lines += ["", "## Ownership", ""]
    corp = [p for p in res if p["corp_or_llc_owned"]]
    out_of_state = [p for p in res if p["owner_out_of_state"]]
    orows = []
    if has_owners(ps):
        elsewhere = [p for p in res if mails_elsewhere(p)]
        orows.append(["Tax bill mailed somewhere other than the property", f"{len(elsewhere):,}",
                      f"{len(elsewhere) / len(res) * 100:.1f}%"])
    orows += [["Tax bill mailed outside Kansas", f"{len(out_of_state):,}", f"{len(out_of_state) / len(res) * 100:.1f}%"],
              ["Owned by a corporation, LLC, or partnership", f"{len(corp):,}", f"{len(corp) / len(res) * 100:.1f}%"]]
    lines += _table(["Residential parcels", "Count", "Share"], orows)
    lines += ["", "**Out-of-state mailing addresses by state:**", ""]
    states = Counter(p["owner_mail_state"] or "(outside the US)" for p in out_of_state)
    lines += _table(["State", "Residential parcels"], [[s, n] for s, n in states.most_common(10)])
    lines += ["", "**Residential parcels owned by a corporation, LLC, or partnership, by use:**", ""]
    lines += _table(["Use", "Parcels", "Share of that use"],
                    [[k, n, f"{n / sum(1 for p in res if p['function'] == k) * 100:.1f}%"]
                     for k, n in Counter(p["function"] for p in corp).most_common(8)])
    lines += ["", "Some owners have their bill sent to a bank, a trust, or a second home. "
              "A mailing address only shows where the bill goes, not where the owner lives. "
              "\"Corporation, LLC, or partnership\" means an owner name includes LLC, L.C., "
              "Inc., Corp., Company, Ltd., LP, LLP, or Partners/Partnership. That includes "
              "homes associations and incorporated nonprofits. It doesn't include trusts or "
              "trust companies acting as trustee."]

    exempt = [p for p in ps if p["class"] == "Exempt"]
    lines += ["", "## Tax-exempt property", "",
              f"{len(exempt):,} parcels, {money(sum(p['total_value'] or 0 for p in exempt))} appraised.", ""]
    lines += _table(["Use", "Parcels", "Appraised"],
                    [[k, f"{g['parcels']:,}", money(g["appraised"])]
                     for k, g in sorted(group_totals(exempt, lambda p: p["function"] or "—"),
                                        key=lambda kv: -kv[1]["appraised"])])

    lines += ["", "## Tax units", "",
              "A tax unit is a set of parcels that all pay the same combination of taxing "
              "jurisdictions. Most of the city is one unit. The smaller units mark special "
              "districts. For example, 0897UT covers the Meadowbrook redevelopment.", ""]
    urows = []
    for k, g in group_totals(ps, lambda p: p["tax_unit"] or "—"):
        examples = Counter(p["address"] for p in ps if p["tax_unit"] == k and p["address"]
                           and not p["address"].startswith("0 ")).most_common(3)
        urows.append([k, f"{g['parcels']:,}", money(g["appraised"]), money(g["assessed"]),
                      ", ".join(a for a, _ in examples)])
    lines += _table(["Tax unit", "Parcels", "Appraised", "Assessed", "Example addresses"], urows, right={1, 2, 3})

    lines += ["", "## Largest parcels by assessed value", ""]
    lines += _table(["Address", "Use", "Assessed", "Appraised"],
                    [[p["address"], p["function"], money(p["assessed_total"]), money(p["total_value"])]
                     for p in sorted(ps, key=lambda p: -(p["assessed_total"] or 0))[:15]], right={2, 3})

    lines += ["", "## Data notes", "",
              "- Source: Johnson County AIMS parcel export. Use `pvc parcels --address ...` to see "
              "one parcel's record and its link to the county's property page.",
              "- Grade, CDU (condition, desirability, usefulness), and condition are the "
              "county's own ratings. They don't always track value. For example, an expensive "
              "house can still be rated \"Very Poor\" on CDU.",
              "- Dwelling units come from the county's land-use coding. The count for a "
              "large apartment complex may not match the true number of units."]
    return "\n".join(lines) + "\n"
